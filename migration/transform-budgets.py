#!/usr/bin/env python3
"""
Transforms a year's 'Verlauf' tab into budgets (spec.md §2.7) plus the
grouping tags its breakdown lines need. Plan0 and Plan1 are imported; Prog
is not — the app computes it live from transactions.

Usage: transform-budgets.py [year] (default 2025), after
transform-transactions.py has run for that year.

Output goes to migration/out/ (gitignored: this repo is public and the
output contains personal data).
"""
import json
import re
import sys
from collections import defaultdict

import openpyxl

YEAR = int(sys.argv[1]) if len(sys.argv) > 1 else 2025
SRC = f"/tmp/claude-0/-home-user-Geld/579a8311-2b64-5844-81b1-64afabfc8004/scratchpad/geld{YEAR}.xlsx"
OUT_BUDGETS = f"/home/user/Geld/migration/out/budgets-{YEAR}.json"
OUT_TAGS = f"/home/user/Geld/migration/out/tags-breakdown-{YEAR}.json"
TRANSACTIONS = f"/home/user/Geld/migration/out/transactions-{YEAR}.json"

# Verlauf layout: D = category, E = Prog/Plan1/Plan0, F = breakdown line (or
# trip), G = sub-line under a trip, J..U = Jan..Dec.
COL_CAT, COL_TYPE, COL_LINE, COL_SUB = 3, 4, 5, 6
MONTHS = range(9, 21)
FIRST_ROW = 40  # the last row is the tab's own last row (grows year to year)
# Categories whose top-level breakdown lines are trips/projects (teal tags).
PROJECT_CATEGORIES = {"sonstiges-urlaube"}
# Blocks whose top line is its own formula and whose breakdown lines are
# stale copies that no longer feed it: the top line is what the sheet
# actually shows, so it is imported alone and the lines are dropped.
STALE_LINES = {(2026, "lebenshaltung-ausstattung-sophia", "plan1")}

ALLOCATION_TAG_MAP = {
    "Für Sparen Familie": "sparen-familie",
    "Für Sparen Sophia": "sparen-sophia",
    "Für Sparen Julia": "sparen-julia",
    "Für Steuern": "ruecklagen-steuern",
    "Für Anlage Familie": "anlage-familie",
    "Für Anlage Sophia": "anlage-sophia",
    "Für Tagesgeld": "tagesgeld",
}


def slug(s):
    s = s.lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def cents(v):
    return round(v * 100) if isinstance(v, (int, float)) else 0


def month_cents(row):
    """The 12 month values in cents, rounded so their sum is still exactly
    the row's yearly total: a yearly figure spread over months (e.g. 1000/12)
    leaves sub-cent parts, and rounding each month on its own would turn
    Markus's -1000.00 into -999.96. Leftover cents go to the months with the
    largest remainders."""
    raw = [row[i] * 100 if isinstance(row[i], (int, float)) else 0 for i in MONTHS]
    out = [round(v) for v in raw]
    short = round(sum(raw)) - sum(out)
    step = 1 if short > 0 else -1
    for i in sorted(range(12), key=lambda i: -(raw[i] - out[i]) * step)[:abs(short)]:
        out[i] += step
    return out


def parse_verlauf(ws, categories):
    """Returns {(targetKind, targetId, plan): {"top": [12 cents], "lines": [(parent, label, [12 cents])]}}."""
    blocks = defaultdict(lambda: {"top": None, "lines": []})
    target, plan, parent = None, None, None
    for row in ws.iter_rows(min_row=FIRST_ROW, max_row=ws.max_row, values_only=True):
        cat, typ, line, sub = row[COL_CAT], row[COL_TYPE], row[COL_LINE], row[COL_SUB]
        values = month_cents(row)
        if cat:
            if cat in categories:
                target = ("expense", categories[cat])
            elif cat in ALLOCATION_TAG_MAP:
                target = ("savings-transfer", ALLOCATION_TAG_MAP[cat])
            else:
                raise ValueError(f"unknown Verlauf row label: {cat!r}")
            plan, parent = None, None
        if typ in ("Plan0", "Plan1"):
            plan, parent = typ.lower(), None
            blocks[(*target, plan)]["top"] = values
            continue
        if typ == "Prog" or plan is None:
            continue
        if isinstance(line, str) and line.strip():
            if isinstance(sub, str) and sub.strip():
                parent = line.strip()
                blocks[(*target, plan)]["lines"].append((parent, sub.strip(), values))
            else:
                parent = None
                blocks[(*target, plan)]["lines"].append((None, line.strip(), values))
        elif isinstance(sub, str) and sub.strip():
            if parent is None:
                raise ValueError(f"sub-line {sub!r} without a parent line in {target}")
            blocks[(*target, plan)]["lines"].append((parent, sub.strip(), values))
    return blocks


def main():
    with open("/home/user/Geld/migration/seed/categories.json") as f:
        categories = {c["name"]: c["id"] for c in json.load(f) if c["parentCategoryId"]}
    ws = openpyxl.load_workbook(SRC, data_only=True)["Verlauf"]
    blocks = parse_verlauf(ws, categories)

    tags, budgets, problems = {}, [], []

    def tag_for(target_id, parent, label):
        if parent:
            pid = f"{target_id}--{slug(parent)}"
            tags.setdefault(pid, {"id": pid, "name": parent, "parentTag": None})
            tid = f"{pid}--{slug(label)}"
            tags.setdefault(tid, {"id": tid, "name": label, "parentTag": pid})
        else:
            tid = f"{target_id}--{slug(label)}"
            tags.setdefault(tid, {"id": tid, "name": label, "parentTag": None})
        grouping = "project" if target_id in PROJECT_CATEGORIES else None
        for t in (tags[tid], tags.get(tags[tid]["parentTag"] or "", {})):
            if t:
                t.update({"class": "grouping", "reconciliationTargetAccountIds": [],
                          "groupingType": grouping, "archived": False})
        return tid

    def emit(kind, target_id, plan, month, amount, breakdown):
        if not amount:
            return
        budgets.append({
            "id": f"b-{YEAR}-{plan}-{target_id}-{breakdown or 'top'}-{month:02d}",
            "year": YEAR, "month": month, "planVersion": plan, "type": kind,
            "categoryId": target_id if kind == "expense" else None,
            "allocationTagId": target_id if kind == "savings-transfer" else None,
            "breakdownTagId": breakdown, "plannedAmountCents": amount, "note": "",
        })

    for (kind, target_id, plan), b in blocks.items():
        if not b["lines"] or (YEAR, target_id, plan) in STALE_LINES:
            for m, v in enumerate(b["top"], start=1):
                emit(kind, target_id, plan, m, v, None)
            continue
        # With breakdown lines, the top line is computed (spec.md §2.7) — only
        # the lines are stored, and they must add up to the sheet's top line.
        for m in range(12):
            s = sum(vals[m] for _, _, vals in b["lines"])
            if s != b["top"][m]:
                problems.append(f"{target_id} {plan} month {m+1}: lines sum {s/100} != top line {b['top'][m]/100}")
        for parent, label, vals in b["lines"]:
            tid = tag_for(target_id, parent, label)
            for m, v in enumerate(vals, start=1):
                emit(kind, target_id, plan, m, v, tid)

    # --- check: group totals per month against the sheet's summary rows ---
    summary = {}
    label = None
    for row in ws.iter_rows(min_row=7, max_row=18, values_only=True):
        if row[COL_CAT]:
            label = row[COL_CAT].strip(" *")
        if row[COL_TYPE] in ("Plan0", "Plan1"):
            summary[(label, row[COL_TYPE].lower())] = month_cents(row)
    with open("/home/user/Geld/migration/seed/categories.json") as f:
        income_ids = {c["id"] for c in json.load(f) if c["parentCategoryId"] == "einnahmen"}
    group_of = lambda b: ("Einnahmen" if b["categoryId"] in income_ids else "Ausgaben") if b["type"] == "expense" \
        else ("Tagesgeld" if b["allocationTagId"] == "tagesgeld" else "Rücklagen")
    computed = defaultdict(lambda: [0] * 12)
    for b in budgets:
        computed[(group_of(b), b["planVersion"])][b["month"] - 1] += b["plannedAmountCents"]
    # Rows are rounded one by one (see month_cents), the sheet's summary
    # rows sum the unrounded values — so a cent or two per month is rounding.
    print("group totals per month vs. Verlauf summary rows:")
    for key, sheet in sorted(summary.items()):
        diffs = [(m + 1, (computed[key][m] - sheet[m]) / 100) for m in range(12) if computed[key][m] != sheet[m]]
        ok = all(abs(d) <= 0.02 for _, d in diffs)
        if not ok:
            problems.append(f"{key}: {[x/100 for x in computed[key]]} vs sheet {[x/100 for x in sheet]}")
        note = f"  (rounding: {diffs})" if diffs and ok else ""
        print(f"  {'OK ' if ok else 'BAD'} {key[0]:10s} {key[1]}  year {sum(computed[key])/100:>12,.2f}{note}")

    # --- check: savings Prog rows against the migrated transactions ---
    # A savings-transfer plan is negative when money is set aside, so its
    # actual is minus the tag's change (spec.md §2.7/§2.8 sign rules).
    import importlib.util
    spec = importlib.util.spec_from_file_location("tx", "/home/user/Geld/migration/transform-transactions.py")
    tx = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(tx)
    with open("/home/user/Geld/migration/seed/tags-allocation.json") as f:
        targets = {t["id"]: set(t["reconciliationTargetAccountIds"]) for t in json.load(f)}
    with open(TRANSACTIONS) as f:
        txs = json.load(f)
    delta = defaultdict(lambda: [0] * 12)
    for t in txs:
        for line in t["lines"]:
            for tag in line["tags"]:
                if tag in targets:
                    delta[tag][int(t["date"][5:7]) - 1] -= tx.allocation_tag_delta(t, line, targets[tag])
    # A year still in progress only has actuals up to its last transaction;
    # the sheet's Prog for later months is a forecast, not comparable.
    last_month = max(int(t["date"][5:7]) for t in txs)
    print(f"savings actuals per month vs. Verlauf Prog rows (months 1-{last_month}):")
    for row in ws.iter_rows(min_row=FIRST_ROW, max_row=ws.max_row, values_only=True):
        if row[COL_TYPE] == "Prog" and row[COL_CAT] in ALLOCATION_TAG_MAP:
            tag = ALLOCATION_TAG_MAP[row[COL_CAT]]
            prog = [cents(row[i]) for i in MONTHS]
            diffs = [(m + 1, (delta[tag][m] - prog[m]) / 100) for m in range(last_month) if delta[tag][m] != prog[m]]
            print(f"  {'OK ' if not diffs else 'DIFF'} {tag:20s} {diffs if diffs else ''}")

    with open(OUT_BUDGETS, "w") as f:
        json.dump(budgets, f, ensure_ascii=False, indent=2)
    with open(OUT_TAGS, "w") as f:
        json.dump(sorted(tags.values(), key=lambda t: t["id"]), f, ensure_ascii=False, indent=2)
    print(f"{len(budgets)} budget rows, {len(tags)} breakdown tags")
    for p in problems:
        print("  PROBLEM:", p)
    print("ALL BUDGET CHECKS PASS" if not problems else f"{len(problems)} PROBLEMS")


if __name__ == "__main__":
    main()
